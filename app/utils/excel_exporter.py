"""Excel export utilities using openpyxl."""
from __future__ import annotations

from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import (
    Font, Alignment, PatternFill, Border, Side, numbers
)
from openpyxl.utils import get_column_letter

from app.utils.formatters import fmt_date, fmt_datetime, fmt_currency


HEADER_FILL = PatternFill("solid", fgColor="1A1A2E")
HEADER_FONT = Font(color="FFD700", bold=True, size=11)
GOLD_FILL = PatternFill("solid", fgColor="FFF8E1")
BORDER = Border(
    bottom=Side(style="thin", color="D0D0D0"),
    right=Side(style="thin", color="D0D0D0"),
)


def _apply_header(ws, headers: list, row: int = 1):
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER


def _auto_width(ws):
    for col in ws.columns:
        max_length = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_length + 4, 40)


def generate_report_xlsx(report_key: str, from_dt, to_dt, path: str):
    from app.services.loan_service import get_loan_service
    from app.config import get_config
    cfg = get_config()

    wb = Workbook()
    ws = wb.active
    ws.title = report_key.replace("_", " ").title()[:31]

    # Business header
    ws.merge_cells("A1:G1")
    ws["A1"] = cfg.business_name
    ws["A1"].font = Font(bold=True, size=14)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.merge_cells("A2:G2")
    ws["A2"] = f"{ws.title}  |  {fmt_date(from_dt)} to {fmt_date(to_dt)}"
    ws["A2"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[3].height = 6

    if report_key == "active_loans":
        headers = ["Loan #", "Customer", "Mobile", "Principal ₹", "Disbursed", "Maturity", "Outstanding ₹"]
        _apply_header(ws, headers, row=4)
        loans = get_loan_service().search_loans(status="active", limit=5000)
        for i, lo in enumerate(loans, start=5):
            ws.cell(i, 1, lo["loan_number"])
            ws.cell(i, 2, lo["customer_name"])
            ws.cell(i, 3, lo["customer_mobile"])
            ws.cell(i, 4, float(lo["principal_amount"] or 0))
            ws.cell(i, 5, str(fmt_date(lo["disbursement_date"])))
            ws.cell(i, 6, str(fmt_date(lo["maturity_date"])))
            ws.cell(i, 7, float(lo["total_outstanding"] or 0))
    elif report_key == "overdue_loans":
        headers = ["Loan #", "Customer", "Mobile", "Overdue Since", "Outstanding ₹"]
        _apply_header(ws, headers, row=4)
        loans = get_loan_service().search_loans(status="overdue", limit=5000)
        for i, lo in enumerate(loans, start=5):
            ws.cell(i, 1, lo["loan_number"])
            ws.cell(i, 2, lo["customer_name"])
            ws.cell(i, 3, lo["customer_mobile"])
            ws.cell(i, 4, str(fmt_date(lo["maturity_date"])))
            ws.cell(i, 5, float(lo["total_outstanding"] or 0))
    else:
        headers = ["Report", "Details"]
        _apply_header(ws, headers, row=4)
        ws.cell(5, 1, report_key)
        ws.cell(5, 2, f"From {fmt_date(from_dt)} to {fmt_date(to_dt)}")

    _auto_width(ws)
    ws.freeze_panes = "A5"

    # Footer
    last_row = ws.max_row + 2
    ws.cell(last_row, 1, f"Generated: {fmt_datetime(datetime.now())}")
    ws.cell(last_row, 1).font = Font(italic=True, color="888888")

    wb.save(path)
